"""Reading resumes/job descriptions from files and basic text utilities."""
import re
from pathlib import Path

STOPWORDS = set("""
a about above after again against all also am an and any are as at be because been before being
below between both but by can could did do does doing down during each etc few for from further
had has have having he her here hers herself him himself his how i if in into is it its itself
just me more most my myself no nor not now of off on once only or other our ours ourselves out
over own same she should so some such than that the their theirs them themselves then there these
they this those through to too under until up very was we were what when where which while who
whom why will with would you your yours yourself yourselves within across per via using use used
able ability strong good great excellent work working experience experienced years year role team
teams including include includes new well must preferred plus required requirements responsibilities
job candidate candidates looking join company knowledge understanding skills skill familiarity
familiar proficiency proficient hands-on minimum qualifications qualification degree bachelor
bachelors master masters related field similar environment opportunity opportunities apply
code coding engineer engineers developer developers associate closely exposure interest grasp solid
clean reliable ship shipping day days various multiple ensure etc help helps like want seeking
ideal ideally preferably across within least tools tool technologies technology plus
""".split())

WORD_RE = re.compile(r"[a-z][a-z0-9+#./-]*[a-z0-9+#]|[a-z]")


class UnsupportedFileError(ValueError):
    pass


def read_text(path):
    """Extract plain text from a .pdf, .docx, .txt or .md file."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        import docx
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix in (".txt", ".md", ""):
        return path.read_text(encoding="utf-8", errors="ignore")
    raise UnsupportedFileError(f"Unsupported file type '{suffix}'. Use PDF, DOCX or TXT.")


def normalize(text):
    text = text.replace("•", "\n").replace("●", "\n").replace("", "\n")
    text = re.sub(r"[–—]", "-", text)
    return text


def tokenize(text):
    """Lower-cased word tokens with stopwords removed and trailing punctuation stripped."""
    return [w.strip("./-") for w in WORD_RE.findall(text.lower())
            if w.strip("./-") and w.strip("./-") not in STOPWORDS and len(w.strip("./-")) > 1]


def phrase_pattern(phrase):
    """Regex that matches a phrase as a whole word, tolerant of symbols like C++ or Node.js."""
    escaped = re.escape(phrase.lower()).replace(r"\ ", r"[\s\-]+")
    return re.compile(rf"(?<![a-z0-9+#.]){escaped}(?![a-z0-9+#])")
